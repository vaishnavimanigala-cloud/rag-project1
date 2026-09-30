from pathlib import Path

try:
    import fitz
except ImportError:  # pragma: no cover
    fitz = None

from docx import Document


def _extract_pdf_text(path):
    if fitz is None:
        raise RuntimeError("PyMuPDF is not installed. Install the project requirements first.")

    pdf_document = fitz.open(str(path))
    pages = []
    parts = []

    for page_number in range(len(pdf_document)):
        page = pdf_document[page_number]
        text = page.get_text("text")
        pages.append({"page": page_number + 1, "text": text})
        parts.append(text)

    pdf_document.close()
    combined_text = "\n\n".join(parts)
    return {
        "source_name": path.name,
        "file_type": "pdf",
        "text": combined_text,
        "pages": pages,
        "page_count": len(pages),
        "char_count": len(combined_text),
    }


def _extract_docx_text(path):
    document = Document(str(path))
    paragraphs = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            paragraphs.append(text)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    paragraphs.append(text)

    combined_text = "\n".join(paragraphs)
    return {
        "source_name": path.name,
        "file_type": "docx",
        "text": combined_text,
        "pages": [{"page": 1, "text": combined_text}],
        "page_count": 1,
        "char_count": len(combined_text),
    }


def extract_document_text(file_path):
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return _extract_pdf_text(path)
    if suffix == ".docx":
        return _extract_docx_text(path)

    raise ValueError(f"Unsupported file type: {suffix}. Please upload a PDF or DOCX file.")
